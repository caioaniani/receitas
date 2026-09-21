"""Preço de MP ainda não informado não vira custo zero nem impede pesagem."""

import pytest

from app.extensions import db
from app.models import MateriaPrima, Produto, ProdutoItem, Receita, ReceitaIngrediente
from app.services.custos import (
    calcular_custo_produto,
    calcular_custos_produtos,
    calcular_custos_receitas,
)


def _receita(nome, *, ingredientes=(), peso=100):
    rec = Receita(nome=nome, peso_base=1000, peso_unitario=peso,
                  rendimento_qtd=1, rendimento_unidade='un')
    rec.ingredientes.extend(ingredientes)
    db.session.add(rec)
    db.session.flush()
    return rec


def _produto(nome, *, tipo, componente, quantidade=1):
    produto = Produto(nome=nome, ativo=True, custo_embalagem=2)
    db.session.add(produto)
    db.session.flush()
    ids = {'receita': 'receita_id', 'mp': 'materia_prima_id',
           'produto': 'produto_componente_id'}
    produto.itens.append(ProdutoItem(
        tipo=tipo, item_nome=componente.nome, quantidade=quantidade,
        **{ids[tipo]: componente.id}))
    db.session.flush()
    return produto


@pytest.mark.parametrize(('tipo', 'qtd', 'unidade', 'peso_mp'), [
    ('mp', 100, 'g', None),
    ('mp_direto', 1000, 'g', None),
    ('mp_un', 10, 'un', 100),
    ('mp_direto', 1000, 'un', 100),
])
def test_custo_pendente_preserva_peso_rendimento_e_nao_vira_ciclo(
        app, tipo, qtd, unidade, peso_mp):
    mp = MateriaPrima(nome='Azeitonas', unidade=unidade,
                      peso_unidade=peso_mp, custo_por_kg=None)
    db.session.add(mp)
    rec = _receita('Receita sem preço de ingrediente', ingredientes=[
        ReceitaIngrediente(tipo=tipo, ingrediente_nome=' aZEITONAS ', porcentagem=qtd)])
    db.session.commit()
    resultado = calcular_custos_receitas()
    assert resultado['custos'][rec.nome] is None
    assert resultado['pesos'][rec.nome] == 100
    assert resultado['circulares'] == []
    assert resultado['custos_pendentes'] == [rec.nome]
    assert resultado['mp_info'][mp.nome]['custo_por_kg'] is None
    fabricado = resultado['fabricados'][0]
    assert fabricado['rendimento'] == 10
    assert fabricado['custo_un'] is None and fabricado['custo_pendente']


def test_zero_explicito_e_ingrediente_sem_consumo_nao_sao_custo_pendente(app):
    db.session.add_all([
        MateriaPrima(nome='Água sem custo', unidade='g', custo_por_kg=0),
        MateriaPrima(nome='MP sem preço não usada', unidade='g', custo_por_kg=None),
    ])
    rec = _receita('Produto de custo zero', ingredientes=[
        ReceitaIngrediente(tipo='mp', ingrediente_nome='Água sem custo', porcentagem=100),
        ReceitaIngrediente(tipo='mp', ingrediente_nome='MP sem preço não usada', porcentagem=0),
    ])
    db.session.commit()
    resultado = calcular_custos_receitas()
    assert resultado['custos'][rec.nome] == 0
    assert resultado['custos_pendentes'] == []
    assert resultado['fabricados'][0]['custo_pendente'] is False


def test_pendencia_propaga_por_sub_receita_retorno_e_cesta_e_some_apos_informar_preco(app):
    mp = MateriaPrima(nome='Azeitonas', unidade='g', custo_por_kg=None)
    db.session.add(mp)
    origem = _receita('Z Origem com azeitonas', ingredientes=[
        ReceitaIngrediente(tipo='mp', ingrediente_nome=mp.nome, porcentagem=100)])
    retorno = _receita('B Retorno sem ficha', peso=None)
    origem.retorno_receita_id = retorno.id
    filho = _receita('A Montagem do retorno', ingredientes=[
        ReceitaIngrediente(tipo='receita', ingrediente_nome='Nome antigo',
                           sub_receita=retorno, porcentagem=2)])
    simples = _produto('Cesta simples', tipo='receita', componente=filho)
    externa = _produto('Cesta de cestas', tipo='produto', componente=simples, quantidade=2)
    direta = _produto('Produto com MP direta', tipo='mp', componente=mp, quantidade=100)
    db.session.commit()

    resultado = calcular_custos_receitas()
    assert set(resultado['custos_pendentes']) == {origem.nome, retorno.nome, filho.nome}
    assert resultado['circulares'] == []
    assert resultado['pesos'][retorno.nome] == 100
    assert next(f for f in resultado['fabricados'] if f['nome'] == filho.nome)['rendimento'] == 2
    produtos = calcular_custos_produtos(resultado['custos'], resultado['mp_info'])
    assert produtos == {simples.nome: None, externa.nome: None, direta.nome: None}
    assert calcular_custo_produto(simples, resultado['custos'], resultado['mp_info']) is None
    assert calcular_custo_produto(direta, resultado['custos'], resultado['mp_info']) is None

    mp.custo_por_kg = 10
    db.session.commit()
    resultado = calcular_custos_receitas()
    assert resultado['custos_pendentes'] == []
    assert resultado['circulares'] == []
    assert all(resultado['custos'][rec.nome] == 1 for rec in (origem, retorno, filho))
    assert calcular_custos_produtos(resultado['custos'], resultado['mp_info']) == {
        simples.nome: 3, externa.nome: 8, direta.nome: 3}


def test_retorno_com_uma_origem_pendente_nao_escolhe_custo_menor_conhecido(app):
    db.session.add_all([
        MateriaPrima(nome='MP conhecida', unidade='g', custo_por_kg=10),
        MateriaPrima(nome='MP pendente', unidade='g', custo_por_kg=None),
    ])
    retorno = _receita('Retorno comum', peso=None)
    for nome, ingrediente in [('Origem com custo', 'MP conhecida'),
                              ('Origem pendente', 'MP pendente')]:
        origem = _receita(nome, ingredientes=[
            ReceitaIngrediente(tipo='mp', ingrediente_nome=ingrediente, porcentagem=100)])
        origem.retorno_receita_id = retorno.id
    db.session.commit()
    resultado = calcular_custos_receitas()
    assert resultado['custos'][retorno.nome] is None
    assert resultado['pesos'][retorno.nome] == 100
    assert retorno.nome not in resultado['circulares']


def _cenario_pendente():
    mp = MateriaPrima(nome='Ingrediente pendente', unidade='g', custo_por_kg=None)
    db.session.add(mp)
    rec = _receita('Pão de custo pendente', ingredientes=[
        ReceitaIngrediente(tipo='mp', ingrediente_nome=mp.nome, porcentagem=100)])
    rec.preco_venda = rec.preco_loja = rec.preco_site = 10
    produto = _produto('Cesta com custo pendente', tipo='receita', componente=rec)
    db.session.commit()
    return mp, rec, produto


def _cliente(app, usuario):
    cliente = app.test_client()
    with cliente.session_transaction() as sessao:
        sessao['_user_id'] = str(usuario.id)
        sessao['_fresh'] = True
    return cliente


def test_leitores_financeiros_mostram_pendencia_sem_margem_ficticia(app, admin_user):
    from app.services import impostos
    from app.services.copilot import _read_consultar_margem
    from app.services.saude_negocio import resumo_receitas

    _mp, rec, _produto_pendente = _cenario_pendente()
    assert impostos.lucro_liquido(10, None, .1) is None
    assert impostos.margem_liquida(10, None, .1) is None
    assert impostos.lucro_liquido(10, 0, .1) == 9
    cliente = _cliente(app, admin_user)
    for url in ['/dashboard', '/rentabilidade', '/relatorios/custos', '/relatorios/ingredientes']:
        resposta = cliente.get(url)
        assert resposta.status_code == 200, url
        assert 'Custo pendente' in resposta.get_data(as_text=True), url
    csv = cliente.get('/relatorios/custos/csv').get_data(as_text=True)
    assert 'Custo pendente,10.00,,10.00,,10.00,' in csv
    categorias = cliente.get('/relatorios/dashboards/api/margem-categoria').get_json()
    assert categorias['valores'] == [None]
    texto = _read_consultar_margem({'nome': rec.nome}, admin_user)['texto']
    assert 'Custo pendente' in texto
    assert 'margem líq.' not in texto
    saude = resumo_receitas()
    assert saude['custos_pendentes'] == [rec.nome]
    assert not saude['margem_critica']


def test_api_e_exportacao_preservam_custo_desconhecido_e_zero_explicito(app):
    from io import BytesIO

    from openpyxl import load_workbook

    from app.services.precos_export import gerar_xlsx_precos

    mp, rec, produto = _cenario_pendente()
    zero = MateriaPrima(nome='MP zero explícito', unidade='un', custo_por_kg=0)
    db.session.add(zero)
    db.session.commit()
    app.config['CLAUDE_API_TOKEN'] = 'teste-custo-pendente'
    resposta = app.test_client().get('/api/claude/custos', headers={
        'Authorization': 'Bearer teste-custo-pendente'})
    assert resposta.status_code == 200
    dados = resposta.get_json()
    assert next(r for r in dados['receitas'] if r['id'] == rec.id)['custo_unitario'] is None
    assert next(p for p in dados['produtos'] if p['id'] == produto.id)['custo'] is None
    mps = {m['id']: m for m in dados['materias_primas']}
    assert mps[mp.id]['custo_por_kg'] is None
    assert mps[zero.id]['custo_unitario'] == 0
    planilha = load_workbook(BytesIO(gerar_xlsx_precos())).active
    linhas = {linha[0]: linha for linha in planilha.iter_rows(min_row=5, values_only=True)}
    assert linhas[rec.nome][1] is None
    assert linhas[produto.nome][1] is None


def test_nf_transferencia_nao_emite_componentes_com_custo_pendente(app, loja, monkeypatch):
    from app.models import PedidoItem, PedidoLoja
    from app.services import tiny_nf_transf
    from app.utils import hoje

    mp, rec, produto = _cenario_pendente()
    pedido = PedidoLoja(loja_id=loja.id, status='em_transporte', data_entrega=hoje())
    db.session.add(pedido)
    db.session.flush()
    pedido.itens.extend([
        PedidoItem(receita_id=rec.id, quantidade=1),
        PedidoItem(produto_id=produto.id, quantidade=1),
        PedidoItem(materia_prima_id=mp.id, quantidade=100),
    ])
    db.session.commit()
    monkeypatch.setattr(tiny_nf_transf, 'sku_transferencia', lambda *args: 'SKU')
    itens, sem_sku, sem_custo = tiny_nf_transf._payload_itens(pedido)
    assert not itens and not sem_sku
    assert set(sem_custo) == {mp.nome + ' (MP)', rec.nome, produto.nome}


def test_perda_operacional_baixa_ingrediente_sem_preco_e_relatorio_informa_pendencia(
        app, admin_user):
    from app.models import Funcionario, MovimentacaoEstoque
    from app.services import perda_producao

    mp, rec, _produto_pendente = _cenario_pendente()
    mp.estoque_atual = 10000
    funcionario = Funcionario(nome='Padeiro de teste', cpf='999.001.001-99',
                              funcao='Padeiro', ativo=True)
    db.session.add(funcionario)
    db.session.commit()
    resultado = perda_producao.registrar(
        rec.id, 2, 'queimou', admin_user.id, fornada=True,
        funcionario_id=funcionario.id)
    assert resultado['perda_id']
    assert mp.estoque_atual < 10000
    assert MovimentacaoEstoque.query.filter_by(materia_prima_id=mp.id).count() > 0
    relatorio = perda_producao.listar()
    assert relatorio['total_qtd'] == 2
    assert relatorio['total_custo'] is None
    assert relatorio['perdas'][0]['custo_total'] is None
    resposta = _cliente(app, admin_user).get('/producao/perdas')
    assert resposta.status_code == 200
    assert 'Custo pendente' in resposta.get_data(as_text=True)
