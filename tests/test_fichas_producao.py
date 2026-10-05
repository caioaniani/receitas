"""Fichas de produção (planilha do dono, 02/10/2026) → etapas do fluxograma.

"cadastrar e usar essas infos no gantt/fluxograma de producao; apagar todos
os registros que o gantt usa hoje e trocar por esses dessa planilha".
Serviço app/services/fichas_producao.py; tela /receitas/fichas-producao;
seed de startup _seed_fichas_producao_2026_10; sonda /api/claude/fichas-producao.
A planilha REAL embarcada em app/seeds_data é a fixture.
"""
import json
import re
from datetime import time, timedelta

import pytest
from sqlalchemy import text

from app.extensions import db
from app.models import (
    AppConfig,
    PlanejamentoItem,
    PlanejamentoProducao,
    Receita,
    ReceitaEtapa,
    ReceitaIngrediente,
    Usuario,
)
from app.models.producao_batelada import PlanejamentoItemBatelada
from app.services import fichas_producao as svc
from app.services.etapas_receita import resumo_parametros
from app.utils import hoje

NOMES_PLANILHA = [
    'Sourdough Tradicional', 'Sourdough Integral', 'Sourdough 7 Grãos',
    'Sourdough Nozes e Azeitonas', 'Pão Francês Fermentado', 'Brioche',
    'Massa para folhar', 'Pão de Forma Integral',
    'Pão de Forma Integral com Grãos', 'Pão de Cranberry',
    'Croissant Tradicional', 'Pain au Chocolat',
]


def _receita(nome, categoria='Pães', **kw):
    r = Receita(nome=nome, categoria=categoria, rendimento_qtd=4,
                rendimento_unidade='un', peso_base=1000.0,
                capacidade_amassadeira_g=50000, **kw)
    db.session.add(r)
    db.session.flush()
    db.session.add(ReceitaIngrediente(receita_id=r.id, tipo='mp',
                                      ingrediente_nome='Farinha', porcentagem=100))
    db.session.commit()
    return r


def _etapa_velha(r, nome='Velha', descricao='texto do padeiro'):
    e = ReceitaEtapa(receita_id=r.id, ordem=0, nome=nome, duracao_min=99,
                     equipamento=None, ativa=True, descricao=descricao)
    db.session.add(e)
    db.session.commit()
    return e


def _login(app, user):
    c = app.test_client()
    c.post('/auth/login', data={'login': user.login, 'senha': '123'},
           follow_redirects=True)
    return c


@pytest.fixture
def lido():
    return svc.ler_planilha(svc.ARQUIVO_SEED)


# ── leitura da planilha real ─────────────────────────────────────────────────

def test_planilha_embarcada_existe_e_tem_as_12_receitas(lido):
    assert svc.ARQUIVO_SEED.exists()
    assert lido['ordem'] == NOMES_PLANILHA
    assert set(lido['fichas']) == set(NOMES_PLANILHA)


def test_sourdough_tradicional_cinco_etapas_na_ordem_da_ficha(lido):
    etapas = lido['fichas']['Sourdough Tradicional']
    assert [(e['nome'], e['duracao_min'], e['equipamento'], e['ativa'])
            for e in etapas] == [
        (svc.NOME_BATIMENTO, 15, 'amassadeira', True),      # V1 14 + V2 1
        (svc.NOME_CAIXA, 85, 'camara_fria', False),         # 30+30+25 (dobras)
        (svc.NOME_MESA, 90, None, False),                   # 1:30
        (svc.NOME_FINAL, 720, 'camara_fria', False),        # CF 12 h
        (svc.NOME_FORNO, 23, 'forno', True),
    ]
    bat = etapas[0]['parametros']
    assert bat['velocidade_1_min'] == 14 and bat['velocidade_2_min'] == 1
    assert bat['temp_ambiente_c'] == 20 and bat['temp_liquido_c'] == 5.4
    assert bat['temp_massa_final_c'] == 23
    assert 'ponto_massa' not in bat                 # em branco na ficha: fora
    assert re.fullmatch(r'1\. Batimento · linha \d+', bat['origem'])
    caixa = etapas[1]['parametros']
    assert caixa['local'] == 'Câmara fria' and caixa['temp_c'] == 21
    assert caixa['dobras_min'] == [30, 30, 25]
    assert caixa['duracao_origem'] == 'ate_ultima_dobra'
    assert 'duracao_h' not in caixa
    assert etapas[2]['parametros']['tempo_como_escrito'] == '1:30'
    final = etapas[3]['parametros']
    assert final['local'] == 'Câmara fria' and final['temp_c'] == 3
    assert final['tempo_como_escrito'] == '12 h'
    assert final['alternativas'] == [{'local': 'Ambiente', 'temp_c': 21,
                                      'tempo_min': 360,
                                      'tempo_como_escrito': '6 h'}]
    forno = etapas[4]['parametros']
    assert forno.pop('origem').startswith('4. Forno · linha ')
    assert forno == {'forno_c': 235, 'teto': 100, 'lastro': 100}


def test_pao_frances_sem_terceira_dobra_sem_mesa_e_cf_tracinho_ignorado(lido):
    etapas = lido['fichas']['Pão Francês Fermentado']
    assert [e['nome'] for e in etapas] == [
        svc.NOME_BATIMENTO, svc.NOME_CAIXA, svc.NOME_FINAL, svc.NOME_FORNO]
    assert etapas[0]['duracao_min'] == 13                 # 8 + 5
    assert etapas[0]['parametros']['temp_liquido_c'] == 5.9
    caixa = etapas[1]
    assert caixa['duracao_min'] == 60 and caixa['equipamento'] is None
    assert caixa['parametros']['local'] == 'Ambiente'
    assert caixa['parametros']['dobras_min'] == [30, 30]
    final = etapas[2]
    assert final['duracao_min'] == 43 and final['equipamento'] is None
    assert final['parametros']['local'] == 'Ambiente'
    assert final['parametros']['temp_c'] == 20
    assert 'alternativas' not in final['parametros']    # a linha CF era '–'
    assert etapas[3]['duracao_min'] == 13


def test_brioche_v2_zero_fica_registrado_como_zero(lido):
    etapas = lido['fichas']['Brioche']
    assert etapas[0]['duracao_min'] == 30
    assert etapas[0]['parametros']['velocidade_2_min'] == 0
    assert etapas[0]['parametros']['temp_liquido_c'] == 20
    assert etapas[0]['parametros']['temp_massa_final_c'] == 27
    assert etapas[1]['parametros']['dobras_min'] == [20, 20]
    assert etapas[1]['duracao_min'] == 40
    assert etapas[2]['duracao_min'] == 30 and etapas[2]['equipamento'] is None
    forno = dict(etapas[3]['parametros'])
    assert forno.pop('origem').startswith('4. Forno · linha ')
    assert forno == {'forno_c': 170, 'teto': 80, 'lastro': 100}


def test_massa_para_folhar_so_batimento_com_liquido_em_texto(lido):
    etapas = lido['fichas']['Massa para folhar']
    assert len(etapas) == 1
    assert etapas[0]['duracao_min'] == 16                  # 5 + 11
    p = etapas[0]['parametros']
    assert p['temp_liquido_txt'] == 'Água 4,5 / Leite 5,2 / Ovos 7,5'
    assert 'temp_liquido_c' not in p
    assert p['temp_massa_final_c'] == 28.6


def test_croissant_24h_literal_e_pain_so_forno(lido):
    cro = lido['fichas']['Croissant Tradicional']
    assert [(e['nome'], e['duracao_min'], e['equipamento']) for e in cro] == [
        (svc.NOME_FINAL, 1440, None), (svc.NOME_FORNO, 18, 'forno')]
    assert cro[0]['parametros']['temp_c'] == 28
    assert cro[0]['parametros']['tempo_como_escrito'] == '24h'
    assert 'local' not in cro[0]['parametros']
    assert cro[1]['parametros']['teto'] == 80 and cro[1]['parametros']['lastro'] == 20
    pain = lido['fichas']['Pain au Chocolat']
    assert [(e['nome'], e['duracao_min']) for e in pain] == [(svc.NOME_FORNO, 18)]


def test_paes_de_forma_e_cranberry_so_forno(lido):
    for nome in ('Pão de Forma Integral', 'Pão de Forma Integral com Grãos',
                 'Pão de Cranberry'):
        etapas = lido['fichas'][nome]
        assert [(e['nome'], e['duracao_min']) for e in etapas] == [(svc.NOME_FORNO, 25)]
        assert etapas[0]['parametros']['forno_c'] == 180


def test_avisos_dizem_o_que_a_ficha_nao_informa(lido):
    # 6 receitas com fermentação em caixa sem "Duração (h)": o fluxograma usa
    # o tempo até a última dobra e avisa — nada é inventado em silêncio.
    caixa = [a for a in lido['avisos'] if 'fermentação em caixa' in a]
    assert len(caixa) == 6
    assert any(a.startswith('Sourdough Tradicional') and '85 min' in a for a in caixa)
    # 12 células amarelas ("leitura incerta" na legenda do dono): 11 com
    # valor em 7 etapas + 1 EM BRANCO (Pain au Chocolat, "Duração (h)")
    incertos = [a for a in lido['avisos'] if 'leitura incerta na ficha' in a]
    assert len(incertos) == 7
    vazias = [a for a in lido['avisos'] if 'célula amarela em branco' in a]
    assert len(vazias) == 1
    assert vazias[0].startswith('Pain au Chocolat') and 'duração' in vazias[0]
    assert any(a.startswith('Croissant Tradicional: Fermentação final') and
               'temperatura, tempo' in a for a in incertos)
    assert any(a.startswith('Pão Francês Fermentado: Forno') and 'tempo' in a
               for a in incertos)
    assert len(lido['avisos']) == 14


def test_celulas_amarelas_viram_incerto_nos_parametros(lido):
    bri = lido['fichas']['Brioche'][0]['parametros']
    assert bri['incerto'] == ['temp_liquido_c', 'temp_massa_final_c']
    cro = lido['fichas']['Croissant Tradicional'][0]['parametros']
    assert cro['incerto'] == ['temp_c', 'tempo_como_escrito']
    pf = lido['fichas']['Pão Francês Fermentado']
    assert pf[2]['parametros']['incerto'] == ['tempo_como_escrito', 'tempo_min']
    assert pf[3]['parametros']['incerto'] == ['tempo_min']
    assert 'incerto' not in lido['fichas']['Sourdough Integral'][0]['parametros']


def test_planilha_invalida_levanta_erro_legivel(tmp_path):
    with pytest.raises(svc.PlanilhaInvalida):
        svc.ler_planilha(b'isto nao e um xlsx')
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = 'Outra coisa'
    caminho = tmp_path / 'x.xlsx'
    wb.save(caminho)
    with pytest.raises(svc.PlanilhaInvalida, match='batimento'):
        svc.ler_planilha(caminho)


def test_minutos_de_texto_e_celulas():
    assert svc._minutos('12 h') == 720
    assert svc._minutos('43 min') == 43
    assert svc._minutos('24h') == 1440
    assert svc._minutos('1:30') == 90
    assert svc._minutos(timedelta(hours=1, minutes=30)) == 90
    assert svc._minutos(time(1, 30)) == 90
    assert svc._minutos(45) == 45
    assert svc._minutos('–') is None and svc._minutos('-') is None
    assert svc._minutos(None) is None
    assert svc._minutos('1:30:00') == 90
    from datetime import datetime
    # 26:00 numa célula h:mm: o openpyxl devolve 1900-01-01 02:00 (bug do
    # ano bissexto de 1900 — base 1899-12-31, não 1899-12-30)
    assert svc._minutos(datetime(1900, 1, 1, 2, 0)) == 26 * 60
    assert svc._minutos(datetime(1900, 1, 2, 0, 0)) == 48 * 60
    assert svc._minutos(datetime(1900, 1, 1, 0, 0)) == 24 * 60
    assert svc._minutos(datetime(1900, 3, 1, 2, 0)) == 61 * 24 * 60 + 120
    assert svc._como_escrito(datetime(1900, 1, 1, 2, 0)) == '26:00'
    assert svc._minutos(datetime(2026, 10, 2, 1, 30)) == 90
    assert svc._horas_min(12) == 720 and svc._horas_min(1.5) == 90
    assert svc._horas_min('12 h') == 720 and svc._horas_min(time(1, 30)) == 90
    assert svc._numero('–') is None and svc._texto('—') is None


# ── rótulo dos parâmetros (fonte única) ──────────────────────────────────────

def test_resumo_parametros_formata_em_portugues(lido):
    sd = lido['fichas']['Sourdough Tradicional']
    assert resumo_parametros(sd[0]['parametros'], 15) == \
        'V1 14 min · V2 1 min · líquido 5,4 °C · ambiente 20 °C · massa ao fim 23 °C'
    assert resumo_parametros(sd[1]['parametros'], 85) == \
        ('Câmara fria · 21 °C · dobras a cada 30, 30 e 25 min · '
         'leitura incerta na ficha: temperatura')
    assert resumo_parametros(sd[2]['parametros'], 90) == ''   # só o tempo
    assert resumo_parametros(sd[3]['parametros'], 720) == \
        ('Câmara fria · 3 °C · alternativa: Ambiente · 21 °C · 6 h · '
         'leitura incerta na ficha: temperatura')
    assert resumo_parametros(sd[4]['parametros'], 23) == '235 °C · teto 100 · lastro 100'
    mf = lido['fichas']['Massa para folhar'][0]['parametros']
    assert resumo_parametros(mf, 16) == (
        'V1 5 min · V2 11 min · líquido: Água 4,5 / Leite 5,2 / Ovos 7,5 · '
        'ambiente 20 °C · massa ao fim 28,6 °C · '
        'leitura incerta na ficha: líquido, massa ao fim')
    pf = lido['fichas']['Pão Francês Fermentado'][1]['parametros']
    assert resumo_parametros(pf, 60) == 'Ambiente · 21 °C · dobras a cada 30 e 30 min'


def test_resumo_parametros_casos_de_borda():
    assert resumo_parametros(None) == ''
    assert resumo_parametros({}) == ''
    assert resumo_parametros('lixo') == ''
    assert resumo_parametros({'chave_desconhecida': 1, 'origem': 'x'}) == ''
    # duração ilegível: o texto da ficha aparece
    assert resumo_parametros({'tempo_como_escrito': '24h', 'temp_c': 28}, 0) == \
        '28 °C · na ficha: 24h'
    assert resumo_parametros({'dobras_min': [20]}) == 'dobra após 20 min'
    # tipo errado nunca derruba o render (o esquema descarta)
    assert resumo_parametros({'dobras_min': 5}) == ''
    assert resumo_parametros({'alternativas': 7, 'forno_c': 235}) == '235 °C'
    assert resumo_parametros({'alternativas': [{'tempo_min': 'x', 'local': 'A'}]}) == \
        'alternativa: A'
    assert resumo_parametros({'forno_c': 'quente', 'teto': True, 'lastro': 80}) == 'lastro 80'


# ── casamento por nome exato ─────────────────────────────────────────────────

def test_casar_por_nome_normalizado_exato(app):
    sd = _receita('Sourdough Tradicional')
    _receita('Mini Sourdough Tradicional')
    pf = _receita('Pao  frances FERMENTADO')          # acento/caixa/espaços
    arquivado = _receita('Brioche')
    arquivado.arquivada_em = hoje()
    db.session.commit()
    _receita('Pão de Cranberry', categoria='Fornadas Especiais')
    _receita('Pão de Cranberry', categoria='Fornadas Especiais')
    fichas = {n: [] for n in ('Sourdough Tradicional', 'Pão Francês Fermentado',
                              'Brioche', 'Pão de Cranberry', 'Pain au Chocolat')}
    c = svc.casar(fichas)
    assert c['Sourdough Tradicional'] == {'receita_id': sd.id,
                                          'receita_nome': 'Sourdough Tradicional'}
    assert c['Pão Francês Fermentado']['receita_id'] == pf.id
    assert c['Brioche'] == {'erro': 'nao_encontrada'}      # só a arquivada
    assert c['Pão de Cranberry']['erro'] == 'ambigua'
    assert len(c['Pão de Cranberry']['candidatos']) == 2
    assert c['Pain au Chocolat'] == {'erro': 'nao_encontrada'}


# ── aplicar: substitui TUDO, com backup e marcador ───────────────────────────

def test_aplicar_substitui_todas_as_etapas_do_sistema_com_backup(app, lido):
    sd = _receita('Sourdough Tradicional')
    bri = _receita('Brioche')
    fora = _receita('Cookie Callebaut', categoria='Doces')   # não está na planilha
    arquivada = _receita('Pão de Cranberry', categoria='Fornadas Especiais')
    arquivada.arquivada_em = hoje()
    _etapa_velha(sd, 'Velha do sourdough', 'passo escrito pelo padeiro')
    _etapa_velha(fora, 'Velha do cookie', 'texto do cookie')
    _etapa_velha(arquivada, 'Velha da arquivada')
    casamento = svc.casar(lido['fichas'])
    resumo = svc.aplicar(lido['fichas'], casamento, usuario_id=7,
                         arquivo='fichas.xlsx', sha256='abc')
    db.session.commit()

    # universo: TODAS as etapas antigas sumiram (inclusive de receita fora da
    # planilha e de arquivada — que nunca recebe as novas: só ativas casam)
    assert ReceitaEtapa.query.filter_by(receita_id=fora.id).count() == 0
    assert ReceitaEtapa.query.filter_by(receita_id=arquivada.id).count() == 0
    assert ReceitaEtapa.query.filter(ReceitaEtapa.nome.like('Velha%')).count() == 0
    db.session.refresh(sd)
    assert [e.nome for e in sd.etapas] == [
        svc.NOME_BATIMENTO, svc.NOME_CAIXA, svc.NOME_MESA, svc.NOME_FINAL,
        svc.NOME_FORNO]
    assert sd.etapas[0].parametros['velocidade_1_min'] == 14     # JSON round-trip
    assert sd.etapas[3].parametros['alternativas'][0]['tempo_min'] == 360
    db.session.refresh(bri)
    assert len(bri.etapas) == 4
    # backup recuperável: as duas etapas antigas, com o texto do padeiro
    bk = json.loads(AppConfig.get(resumo['backup']))
    nomes_bk = {e['nome']: e for e in bk['etapas']}
    assert nomes_bk['Velha do sourdough']['descricao'] == 'passo escrito pelo padeiro'
    assert nomes_bk['Velha do cookie']['receita_id'] == fora.id
    # marcador com o resumo
    marc = json.loads(AppConfig.get(svc.MARCADOR_IMPORT))
    assert marc['receitas_aplicadas'] == 2 and marc['etapas'] == 9
    assert marc['etapas_anteriores'] == 3 and marc['usuario_id'] == 7
    assert marc['arquivo'] == 'fichas.xlsx' and marc['sha256'] == 'abc'
    assert set(marc['nao_encontradas']) == set(NOMES_PLANILHA) - {
        'Sourdough Tradicional', 'Brioche'}
    assert marc['ambiguas'] == []
    assert svc.ultima_importacao()['backup'] == resumo['backup']


def test_aplicar_reescreve_o_processo_das_ordens_abertas_e_so_delas(app, lido):
    sd = _receita('Sourdough Tradicional')
    _etapa_velha(sd, 'Amassamento velho')
    processo_velho = [{'nome': 'Amassamento velho', 'duracao': '99 min',
                       'duracao_min': 99, 'equipamento': None, 'ativa': True,
                       'descricao': None}]

    def _ordem(dia, produzido=0, alvo=8, **kw):
        pl = PlanejamentoProducao(data=dia, origem='cronograma')
        db.session.add(pl)
        db.session.flush()
        it = PlanejamentoItem(planejamento_id=pl.id, receita_id=sd.id,
                              multiplicador=1, qtd_alvo=alvo,
                              produzido_qtd=produzido, **kw)
        db.session.add(it)
        db.session.flush()
        snap = PlanejamentoItemBatelada(
            item_id=it.id, bateladas=2,
            dados={'unidades': 4, 'mp': [{'id': 1, 'quantidade': 500.0}],
                   'processo': processo_velho})
        db.session.add(snap)
        db.session.commit()
        return it

    hj = hoje()
    aberta = _ordem(hj)
    continuacao = _ordem(hj - timedelta(days=3))          # ainda na janela
    antiga = _ordem(hj - timedelta(days=4))               # fora da janela
    concluida = _ordem(hj, produzido=8)
    dispensada = _ordem(hj, dispensada_em=db.func.now())
    encerrada = _ordem(hj, falta_encerrada_em=db.func.now())

    resumo = svc.aplicar(lido['fichas'], svc.casar(lido['fichas']))
    db.session.commit()
    assert resumo['snapshots_reescritos'] == 2

    def _proc(it):
        snap = db.session.get(PlanejamentoItemBatelada, it.id)
        db.session.refresh(snap)
        return snap.dados['processo'], snap.dados

    proc, dados = _proc(aberta)
    assert [p['nome'] for p in proc] == [
        svc.NOME_BATIMENTO, svc.NOME_CAIXA, svc.NOME_MESA, svc.NOME_FINAL,
        svc.NOME_FORNO]
    assert proc[0]['parametros_label'].startswith('V1 14 min')
    assert dados['mp'] == [{'id': 1, 'quantidade': 500.0}]   # pesagem intacta
    assert dados['unidades'] == 4
    assert _proc(continuacao)[0][0]['nome'] == svc.NOME_BATIMENTO
    for it in (antiga, concluida, dispensada, encerrada):
        assert _proc(it)[0] == processo_velho


# ── Gantt e TV usam as etapas e os parâmetros novos ──────────────────────────

def _plano_enviado(dia, rec, alvo=8):
    pl = PlanejamentoProducao(data=dia, origem='cronograma')
    db.session.add(pl)
    db.session.flush()
    db.session.add(PlanejamentoItem(planejamento_id=pl.id, receita_id=rec.id,
                                    multiplicador=1, qtd_alvo=alvo,
                                    produzido_qtd=0))
    db.session.commit()
    return pl


def test_gantt_mostra_parametros_no_passo_no_destino_e_na_continuacao(app, lido):
    from app.services.gantt import montar_gantt
    sd = _receita('Sourdough Tradicional', dias_producao=1)
    svc.aplicar(lido['fichas'], svc.casar(lido['fichas']))
    db.session.commit()
    hj = hoje()
    _plano_enviado(hj, sd)
    g = montar_gantt(hj)
    prod = next(p for p in g['produtos'] if p['nome'] == 'Sourdough Tradicional')
    etapas = [t['etapa'] for t in prod['tarefas']]
    assert etapas == [svc.NOME_BATIMENTO, svc.NOME_CAIXA, svc.NOME_MESA]
    assert prod['tarefas'][0]['parametros_label'].startswith('V1 14 min · V2 1 min')
    assert prod['tarefas'][1]['parametros_label'] == \
        ('Câmara fria · 21 °C · dobras a cada 30, 30 e 25 min · '
         'leitura incerta na ficha: temperatura')
    assert prod['destino_etapa'] == svc.NOME_FINAL
    assert prod['destino_parametros_label'] == \
        ('Câmara fria · 3 °C · alternativa: Ambiente · 21 °C · 6 h · '
         'leitura incerta na ficha: temperatura')
    # continuação no dia seguinte: só o forno, com os parâmetros do forno
    g2 = montar_gantt(hj + timedelta(days=1))
    cont = next(p for p in g2['produtos'] if p['tipo'] == 'continuacao')
    assert [t['etapa'] for t in cont['tarefas']] == [svc.NOME_FORNO]
    assert cont['tarefas'][0]['parametros_label'] == '235 °C · teto 100 · lastro 100'


def test_rota_gantt_renderiza_parametros(app, admin_user, lido):
    sd = _receita('Sourdough Tradicional', dias_producao=1)
    svc.aplicar(lido['fichas'], svc.casar(lido['fichas']))
    db.session.commit()
    _plano_enviado(hoje(), sd)
    c = _login(app, admin_user)
    html = c.get('/padeiro/gantt').get_data(as_text=True)
    assert 'class="step-params"' in html
    assert 'V1 14 min · V2 1 min · líquido 5,4 °C' in html
    assert 'Câmara fria · 3 °C · alternativa: Ambiente · 21 °C · 6 h · ' \
           'leitura incerta na ficha: temperatura' in html
    assert 'padeiro-sequencia.css?v=20261002a' in html


def test_tv_mise_en_place_traz_parametros_da_ficha_e_do_snapshot(app, admin_user, lido):
    from app.services.bateladas_paes import processo_da_receita
    sd = _receita('Sourdough Tradicional')
    svc.aplicar(lido['fichas'], svc.casar(lido['fichas']))
    db.session.commit()
    c = _login(app, admin_user)
    d = c.get(f'/padeiro/receita/{sd.id}.json?unidades=8').get_json()
    assert d['processo'][0]['parametros_label'].startswith('V1 14 min')
    forno = dict(d['processo'][4]['parametros'])
    assert forno.pop('origem').startswith('4. Forno · linha ')
    assert forno == {'forno_c': 235, 'teto': 100, 'lastro': 100}
    # mesma forma do processo congelado (fonte única)
    assert [p['nome'] for p in processo_da_receita(sd)] == \
        [p['nome'] for p in d['processo']]


# ── seed de startup ──────────────────────────────────────────────────────────

def test_seed_de_startup_aplica_uma_vez_com_contagens_no_marker(app):
    from app.migrations_legacy import _seed_fichas_producao_2026_10
    sd = _receita('Sourdough Tradicional')
    _receita('Brioche')
    fora = _receita('Cookie', categoria='Doces')
    _etapa_velha(fora)
    _seed_fichas_producao_2026_10(app)
    marker = AppConfig.get('seed_fichas_producao_2026_10')
    assert marker.startswith('receitas=2 etapas=9 anteriores=1 snapshots=0 '
                             'nao_encontradas=10 ambiguas=0 sha256=4b99b0bb17ee')
    db.session.refresh(sd)
    assert len(sd.etapas) == 5
    assert ReceitaEtapa.query.filter_by(receita_id=fora.id).count() == 0
    assert svc.ultima_importacao()['origem'] == 'seed'
    # segunda chamada: no-op (o dono pode ter editado depois)
    _etapa_velha(fora, 'Editada pelo dono')
    _seed_fichas_producao_2026_10(app)
    assert ReceitaEtapa.query.filter_by(receita_id=fora.id).count() == 1
    assert AppConfig.get('seed_fichas_producao_2026_10') == marker


# ── tela /receitas/fichas-producao ───────────────────────────────────────────

def _upload(c, url='/receitas/fichas-producao/previa'):
    import io
    dados = svc.ARQUIVO_SEED.read_bytes()
    return c.post(url, data={'planilha': (io.BytesIO(dados), 'fichas.xlsx')},
                  content_type='multipart/form-data')


def test_tela_estado_lista_receitas_e_as_sem_etapas(app, admin_user):
    sd = _receita('Sourdough Tradicional')
    db.session.add(ReceitaEtapa(receita_id=sd.id, ordem=0, nome='Forno',
                                duracao_min=23, equipamento='forno',
                                parametros={'forno_c': 235, 'teto': 100}))
    _receita('Cookie', categoria='Doces')
    db.session.commit()
    c = _login(app, admin_user)
    html = c.get('/receitas/fichas-producao').get_data(as_text=True)
    assert 'Sourdough Tradicional' in html
    assert '⚙ 235 °C · teto 100' in html
    assert '1 receita(s) sem etapas' in html and 'Cookie' in html
    assert 'Nenhuma importação registrada' in html


def test_previa_mostra_casamento_e_nao_grava(app, admin_user):
    """Admin vê a prévia completa mas não o botão de aplicar (a troca é
    global — só o dono)."""
    sd = _receita('Sourdough Tradicional')
    _etapa_velha(sd)
    c = _login(app, admin_user)
    resp = _upload(c)
    assert resp.status_code == 200
    html = resp.get_data(as_text=True)
    assert 'Prévia' in html
    assert f'#{sd.id} Sourdough Tradicional' in html
    assert html.count('não encontrada') == 11
    assert 'V1 14 min · V2 1 min' in html
    assert 'Só o proprietário aplica' in html
    assert 'name="token"' not in html
    assert ReceitaEtapa.query.filter_by(receita_id=sd.id, nome='Velha').count() == 1


def test_dono_ve_o_token_na_previa_e_aplica_pela_tela(app, owner_user):
    import re
    sd = _receita('Sourdough Tradicional')
    _etapa_velha(sd)
    c = _login(app, owner_user)
    html = _upload(c).get_data(as_text=True)
    assert 'Só o proprietário aplica' not in html
    token = re.search(r'name="token" value="([^"]+)"', html).group(1)
    resp = c.post('/receitas/fichas-producao/aplicar', data={'token': token},
                  follow_redirects=True)
    assert resp.status_code == 200
    body = resp.get_data(as_text=True)
    assert 'Fichas aplicadas: 1 receita(s), 5 etapa(s)' in body
    assert '1 etapa(s) anteriores guardadas' in body
    db.session.refresh(sd)
    assert [e.nome for e in sd.etapas][0] == svc.NOME_BATIMENTO
    marc = svc.ultima_importacao()
    assert marc['origem'] == 'tela' and marc['usuario_id'] == owner_user.id
    assert marc['arquivo'] == 'fichas.xlsx'
    assert 'Última importação' in body


def test_aplicar_pela_tela_exige_dono(app, admin_user):
    """Admin comum com um token válido na mão: 403 e nada gravado."""
    from app.blueprints.receitas.routes import _fichas_serializer
    sd = _receita('Sourdough Tradicional')
    _etapa_velha(sd)
    lido = svc.ler_planilha(svc.ARQUIVO_SEED)
    token = _fichas_serializer().dumps({'fichas': lido['fichas'],
                                        'arquivo': 'x.xlsx', 'sha256': 'abc'})
    c = _login(app, admin_user)
    resp = c.post('/receitas/fichas-producao/aplicar', data={'token': token})
    assert resp.status_code == 403
    assert ReceitaEtapa.query.filter_by(receita_id=sd.id, nome='Velha').count() == 1


def test_aplicar_com_token_invalido_ou_vencido_nao_grava(app, owner_user):
    sd = _receita('Sourdough Tradicional')
    _etapa_velha(sd)
    c = _login(app, owner_user)
    resp = c.post('/receitas/fichas-producao/aplicar', data={'token': 'lixo'},
                  follow_redirects=True)
    assert 'prévia expirou' in resp.get_data(as_text=True)
    assert ReceitaEtapa.query.filter_by(receita_id=sd.id).count() == 1


def test_previa_sem_arquivo_ou_invalido_avisa(app, admin_user):
    import io
    c = _login(app, admin_user)
    resp = c.post('/receitas/fichas-producao/previa', data={},
                  follow_redirects=True)
    assert 'Envie a planilha' in resp.get_data(as_text=True)
    resp = c.post('/receitas/fichas-producao/previa',
                  data={'planilha': (io.BytesIO(b'nao e xlsx'), 'x.xlsx')},
                  content_type='multipart/form-data', follow_redirects=True)
    assert 'Não foi possível abrir a planilha' in resp.get_data(as_text=True)


def test_tela_exige_admin(app):
    u = Usuario(nome='func', login='funcfp', papel='funcionario')
    u.set_senha('123')
    db.session.add(u)
    db.session.commit()
    c = _login(app, u)
    assert c.get('/receitas/fichas-producao').status_code == 403
    assert _upload(c).status_code == 403


def test_menu_do_catalogo_tem_o_link(app, admin_user):
    c = _login(app, admin_user)
    html = c.get('/area/catalogo').get_data(as_text=True)
    assert '/receitas/fichas-producao' in html


# ── sonda read-only ──────────────────────────────────────────────────────────

TOKEN = 'token-de-teste-bem-longo-123'


def test_sonda_fichas_producao_e_receita_expoem_etapas(app, lido):
    app.config['CLAUDE_API_TOKEN'] = TOKEN
    sd = _receita('Sourdough Tradicional')
    _receita('Cookie', categoria='Doces')
    svc.aplicar(lido['fichas'], svc.casar(lido['fichas']))
    db.session.commit()
    c = app.test_client()
    h = {'Authorization': 'Bearer ' + TOKEN}
    d = c.get('/api/claude/fichas-producao', headers=h).get_json()
    assert d['ok'] and d['total_receitas'] == 2 and d['com_etapas'] == 1
    assert [r['nome'] for r in d['sem_etapas']] == ['Cookie']
    assert d['importacao']['receitas_aplicadas'] == 1
    rec = d['receitas'][0]
    assert rec['nome'] == 'Sourdough Tradicional' and rec['n_etapas'] == 5
    assert rec['etapas'][4]['parametros_label'] == '235 °C · teto 100 · lastro 100'
    d2 = c.get(f'/api/claude/receita?id={sd.id}', headers=h).get_json()
    assert [e['nome'] for e in d2['receita']['etapas']][0] == svc.NOME_BATIMENTO
    assert d2['receita']['etapas'][0]['parametros']['velocidade_1_min'] == 14
    assert c.get('/api/claude/fichas-producao').status_code == 401


def test_duplicar_receita_copia_os_parametros(app, admin_user):
    sd = _receita('Sourdough Tradicional')
    db.session.add(ReceitaEtapa(receita_id=sd.id, ordem=0, nome='Forno',
                                duracao_min=23, equipamento='forno',
                                parametros={'forno_c': 235, 'teto': 100}))
    db.session.commit()
    c = _login(app, admin_user)
    c.post(f'/receitas/{sd.id}/duplicar', follow_redirects=True)
    copia = Receita.query.filter_by(nome='Cópia de Sourdough Tradicional').first()
    assert copia is not None
    assert copia.etapas[0].parametros == {'forno_c': 235, 'teto': 100}


# ── revisão independente (02/10/2026): casos de borda com planilha sintética ──

def _planilha(tmp_path, abas):
    """xlsx com os cabeçalhos REAIS da planilha do dono; `abas` =
    {'batimento': [linhas], 'caixa': [...], 'final': [...], 'forno': [...]}
    (cada linha uma lista de valores a partir da coluna Receita; valores
    (valor, 'incerto') marcam a célula em amarelo)."""
    import openpyxl
    from openpyxl.styles import PatternFill
    cabecalhos = {
        'batimento': ('1. Batimento', ['Receita', 'Velocidade 1 (min)',
                      'Velocidade 2 (min)', 'Ponto da massa',
                      'Temperatura ambiente (°C)', 'Temperatura do líquido (°C)',
                      'Temperatura após o fim do batimento (°C)']),
        'caixa': ('2. Fermentação', ['Receita', 'Local', 'Temperatura (°C)',
                  'Duração (h)', 'Dobra 1 (min)', 'Dobra 2 (min)', 'Dobra 3 (min)',
                  'MESA – corte à modelagem: tempo final (h:mm)']),
        'final': ('3. Fermentação final', ['Receita',
                  'Local (CF = Câmara fria · A = Ambiente)', 'Temperatura (°C)',
                  'Tempo (como escrito)', 'Tempo (min)']),
        'forno': ('4. Forno', ['Receita', 'Forno (°C)', 'Tempo (min)',
                  'Teto (valor)', 'Lastro (valor)', 'Unidade teto / lastro',
                  'Alvo interno (°C)']),
    }
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    amarelo = PatternFill('solid', fgColor='FFFFF2CC')
    for chave, (titulo, cab) in cabecalhos.items():
        ws = wb.create_sheet(titulo)
        ws.append([titulo])
        ws.append(['nota'])
        linhas = abas.get(chave, [])
        if isinstance(linhas, dict):            # cabeçalho customizado
            cab = linhas['cabecalho']
            linhas = linhas['linhas']
        ws.append(cab)
        for linha in linhas:
            valores = []
            marcados = []
            for i, v in enumerate(linha):
                if isinstance(v, tuple) and len(v) == 2 and v[1] == 'incerto':
                    valores.append(v[0]); marcados.append(i)
                else:
                    valores.append(v)
            ws.append(valores)
            for i in marcados:
                ws.cell(row=ws.max_row, column=i + 1).fill = amarelo
        ws.append(['Legenda: azul = escrito à mão'])
    caminho = tmp_path / 'sintetica.xlsx'
    wb.save(caminho)
    return caminho


def test_grafia_diferente_entre_abas_e_a_mesma_ficha_na_ordem_certa(app, tmp_path):
    caminho = _planilha(tmp_path, {
        'batimento': [['Sourdough Tradicional', 14, 1, None, 20, 5.4, 23]],
        'caixa': [['Sourdough  Tradicional', 'Câmara fria', 21, None, 30, 30, 25,
                   timedelta(minutes=90)]],
        'final': [['sourdough tradicional', 'CF', 3, '12 h', 720]],
        'forno': [['SOURDOUGH TRADICIONAL', 235, 23, 100, 100, None, None]],
    })
    lido = svc.ler_planilha(caminho)
    assert lido['ordem'] == ['Sourdough Tradicional']
    assert [e['nome'] for e in lido['fichas']['Sourdough Tradicional']] == [
        svc.NOME_BATIMENTO, svc.NOME_CAIXA, svc.NOME_MESA, svc.NOME_FINAL,
        svc.NOME_FORNO]
    grafias = [a for a in lido['avisos'] if 'grafia diferente' in a]
    assert len(grafias) == 3
    sd = _receita('Sourdough Tradicional')
    resumo = svc.aplicar(lido['fichas'], svc.casar(lido['fichas']))
    db.session.commit()
    assert resumo['receitas_aplicadas'] == 1
    db.session.refresh(sd)
    assert [(e.ordem, e.nome) for e in sd.etapas] == [
        (0, svc.NOME_BATIMENTO), (1, svc.NOME_CAIXA), (2, svc.NOME_MESA),
        (3, svc.NOME_FINAL), (4, svc.NOME_FORNO)]


def test_ordem_das_colunas_nao_importa_e_coluna_obrigatoria_ausente_avisa(tmp_path):
    caminho = _planilha(tmp_path, {
        'forno': {'cabecalho': ['Receita', 'Unidade teto / lastro',
                                'Tempo (como escrito)', 'Teto (valor)',
                                'Lastro (valor)', 'Tempo (min)', 'Forno (°C)'],
                  'linhas': [['Brioche', '%', '45 min', 80, 100, 45, 170]]},
        'batimento': [['Brioche', 30, 0, None, 20, 20, 27]],
    })
    lido = svc.ler_planilha(caminho)
    forno = lido['fichas']['Brioche'][1]
    assert forno['duracao_min'] == 45
    assert forno['parametros']['teto'] == 80 and forno['parametros']['lastro'] == 100
    assert forno['parametros']['forno_c'] == 170
    assert forno['parametros']['unidade_teto_lastro'] == '%'
    assert resumo_parametros(forno['parametros'], 45) == \
        '170 °C · teto 80 · lastro 100 · (%)'
    # cabeçalho obrigatório ausente: erro legível, não etapa calada
    sem = _planilha(tmp_path, {
        'forno': {'cabecalho': ['Receita', 'Forno (°C)'], 'linhas': [['Brioche', 170]]}})
    with pytest.raises(svc.PlanilhaInvalida, match='tempo'):
        svc.ler_planilha(sem)


def test_traco_e_linha_so_com_temperatura_nao_viram_etapa_principal(tmp_path):
    caminho = _planilha(tmp_path, {
        'final': [['Pão Francês Fermentado', 'CF', '–', '–', None],
                  ['Pão Francês Fermentado', 'A', 20, '43 min', 43],
                  ['Brioche', 'CF', 4, None, None],
                  ['Brioche', 'A', 20, '30 min', 30]],
    })
    lido = svc.ler_planilha(caminho)
    pf = lido['fichas']['Pão Francês Fermentado'][0]
    assert pf['duracao_min'] == 43 and pf['equipamento'] is None
    assert pf['parametros']['local'] == 'Ambiente'
    assert 'alternativas' not in pf['parametros']
    bri = lido['fichas']['Brioche'][0]
    assert bri['duracao_min'] == 30 and bri['parametros']['local'] == 'Ambiente'
    assert bri['parametros']['alternativas'] == [{'local': 'Câmara fria', 'temp_c': 4}]
    assert any('Brioche: fermentação final sem tempo na ficha' in a
               for a in lido['avisos'])
    assert not any('None' in a for a in lido['avisos'])


def test_duracao_em_horas_mesa_numerica_e_divergencia_min_x_texto(tmp_path):
    caminho = _planilha(tmp_path, {
        'caixa': [['Pão Francês Fermentado', 'Ambiente', 21, time(1, 30), 30, 30, None, None],
                  ['Brioche', 'Ambiente', 21, '2 h', 20, 20, None, 1.5],
                  ['Sourdough Integral', 'Ambiente', 21, 'duas horas', 30, 30, 30, None],
                  ['Pão de Cranberry', 'Ambiente', 21, None, None, None, None, None]],
        'final': [['Brioche', 'A', 20, '30 min', 300]],
        'forno': [['Brioche', 170, None, 80, 100, None, None],
                  ['Pão de Cranberry', 180, 25, 80, 100, None, 92]],
    })
    lido = svc.ler_planilha(caminho)
    pf = lido['fichas']['Pão Francês Fermentado'][0]
    assert pf['duracao_min'] == 90 and pf['parametros']['duracao_origem'] == 'ficha'
    assert pf['parametros']['duracao_h'] == 1.5
    bri = lido['fichas']['Brioche']
    assert bri[0]['duracao_min'] == 120                    # '2 h' como texto
    assert bri[1]['nome'] == svc.NOME_MESA and bri[1]['duracao_min'] == 90   # 1.5 h
    assert bri[2]['duracao_min'] == 300                    # vale o número…
    assert any('diferente do escrito ("30 min" = 30 min)' in a for a in lido['avisos'])
    assert bri[3]['duracao_min'] == 0                      # forno sem tempo
    assert any('Brioche: forno sem tempo na ficha' in a for a in lido['avisos'])
    si = lido['fichas']['Sourdough Integral'][0]
    assert si['duracao_min'] == 90 and si['parametros']['duracao_origem'] == 'ate_ultima_dobra'
    assert any('"Duração (h)" ilegível ("duas horas"' in a for a in lido['avisos'])
    cran = lido['fichas']['Pão de Cranberry']
    assert cran[0]['duracao_min'] == 0 and cran[0]['parametros']['duracao_origem'] == 'ausente'
    assert any('sem tempo nem dobras' in a for a in lido['avisos'])
    assert resumo_parametros(cran[1]['parametros'], 25) == \
        '180 °C · teto 80 · lastro 100 · alvo interno 92 °C'


def test_linha_repetida_e_linha_vazia_nao_criam_etapa(tmp_path):
    caminho = _planilha(tmp_path, {
        'batimento': [['Brioche', 30, 0, None, 20, 20, 27],
                      ['Brioche', 31, 0, None, None, None, None],
                      ['Fantasma', None, None, None, None, None, None]],
        'caixa': [['Fantasma', None, None, None, None, None, None, None]],
    })
    lido = svc.ler_planilha(caminho)
    assert lido['ordem'] == ['Brioche']
    assert [e['duracao_min'] for e in lido['fichas']['Brioche']] == [30]
    assert any('Brioche: linha repetida' in a for a in lido['avisos'])


def test_duracao_fora_da_faixa_vira_aviso_e_zero_e_a_regua_e_unica(tmp_path):
    from app.services.etapas_receita import DURACAO_MAX, duracao_valida
    caminho = _planilha(tmp_path, {
        'batimento': [['Brioche', 3000000000, 0, None, None, None, None],
                      ['Sourdough Integral', -5, 2, None, None, None, None]],
    })
    lido = svc.ler_planilha(caminho)
    assert lido['fichas']['Brioche'][0]['duracao_min'] == 0
    assert lido['fichas']['Sourdough Integral'][0]['duracao_min'] == 0
    assert sum('fora da faixa' in a for a in lido['avisos']) == 2
    assert duracao_valida(-1) == 0 and duracao_valida(10 ** 12) == DURACAO_MAX
    assert duracao_valida('abc') == 0 and duracao_valida('15') == 15


def test_celula_amarela_sintetica_marca_incerto(tmp_path):
    caminho = _planilha(tmp_path, {
        'forno': [['Brioche', (170, 'incerto'), (45, 'incerto'), 80, 100, None, None]],
    })
    lido = svc.ler_planilha(caminho)
    p = lido['fichas']['Brioche'][0]['parametros']
    assert p['incerto'] == ['forno_c', 'tempo_min']
    assert resumo_parametros(p, 45) == \
        '170 °C · teto 80 · lastro 100 · leitura incerta na ficha: forno, tempo'


def test_hora_acima_de_24h_le_o_que_o_openpyxl_devolve_de_verdade(tmp_path):
    """Round-trip REAL: célula h:mm com 26:00 / 48:00 / 24:00 / 1:30 gravada
    e relida pelo openpyxl (o datetime de 1900 não é inventado no teste)."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    for linha, horas in enumerate((26, 48, 24, 1.5), start=1):
        c = ws.cell(row=linha, column=1, value=horas / 24.0)
        c.number_format = 'h:mm'
    caminho = tmp_path / 'horas.xlsx'
    wb.save(caminho)
    ws = openpyxl.load_workbook(caminho, data_only=True).active
    lidos = [ws.cell(row=i, column=1).value for i in range(1, 5)]
    assert [svc._minutos(v) for v in lidos] == [1560, 2880, 1440, 90]
    assert [svc._como_escrito(v) for v in lidos] == \
        ['26:00', '48:00', '24:00', '1:30']
    # e a MESA (h:mm) de uma linha real passa pelo mesmo caminho
    caminho2 = _planilha(tmp_path, {
        'caixa': [['Brioche', 'Ambiente', 21, None, 20, 20, None, lidos[0]]],
    })
    lido = svc.ler_planilha(caminho2)
    mesa = lido['fichas']['Brioche'][1]
    assert mesa['nome'] == svc.NOME_MESA and mesa['duracao_min'] == 1560
    assert mesa['parametros']['tempo_como_escrito'] == '26:00'


def test_fermentacao_final_com_o_mesmo_metodo_repetido_vale_a_primeira(tmp_path):
    caminho = _planilha(tmp_path, {
        'final': [['Brioche', 'A', 20, '30 min', 30],
                  ['Brioche', 'A', 24, '45 min', 45],
                  ['Brioche', 'CF', 4, '10 h', 600]],
    })
    lido = svc.ler_planilha(caminho)
    etapa = lido['fichas']['Brioche'][0]
    assert etapa['duracao_min'] == 30 and etapa['parametros']['temp_c'] == 20
    alts = etapa['parametros']['alternativas']
    assert [a['local'] for a in alts] == ['Câmara fria']       # a 2ª 'A' caiu
    repetidas = [a for a in lido['avisos'] if 'linha repetida' in a]
    assert len(repetidas) == 1 and 'método "Ambiente"' in repetidas[0]


def test_celula_amarela_na_alternativa_marca_e_avisa(tmp_path):
    caminho = _planilha(tmp_path, {
        'final': [['Brioche', 'A', 20, '30 min', 30],
                  ['Brioche', 'CF', (4, 'incerto'), '10 h', 600]],
    })
    lido = svc.ler_planilha(caminho)
    p = lido['fichas']['Brioche'][0]['parametros']
    assert p['incerto'] == ['alternativas']
    assert resumo_parametros(p, 30).endswith(
        'alternativa: Câmara fria · 4 °C · 10 h · '
        'leitura incerta na ficha: alternativa')
    avisos = [a for a in lido['avisos'] if 'leitura incerta' in a]
    assert len(avisos) == 1
    assert avisos[0].startswith('Brioche: alternativa da fermentação final '
                                '("Câmara fria")')
    assert 'temperatura' in avisos[0] and 'linha 5' in avisos[0]   # a linha DELA


def test_qualquer_dobra_amarela_marca_as_dobras(tmp_path):
    caminho = _planilha(tmp_path, {
        'caixa': [['Brioche', 'Ambiente', 21, None, 20, (20, 'incerto'), None,
                   None],
                  ['Sourdough Integral', 'Câmara fria', 21, None, 30,
                   (30, 'incerto'), (30, 'incerto'), None]],
    })
    lido = svc.ler_planilha(caminho)
    assert lido['fichas']['Brioche'][0]['parametros']['incerto'] == ['dobras_min']
    # duas dobras amarelas = a chave entra UMA vez
    assert lido['fichas']['Sourdough Integral'][0]['parametros']['incerto'] == \
        ['dobras_min']
    assert sum('célula amarela: dobras' in a for a in lido['avisos']) == 2


def test_local_sozinho_nao_cria_fermentacao_em_caixa(tmp_path):
    caminho = _planilha(tmp_path, {
        'caixa': [['Brioche', 'Ambiente', None, None, None, None, None, None],
                  ['Sourdough Integral', None, None, None, None, None, None,
                   timedelta(minutes=90)],
                  ['Pão de Cranberry', 'Ambiente', None, None, None, None, None,
                   timedelta(minutes=45)]],
        'forno': [['Brioche', 170, 45, None, None, None, None]],
    })
    lido = svc.ler_planilha(caminho)
    # "Local" sozinho: nenhuma etapa de caixa (e aviso)
    assert [e['nome'] for e in lido['fichas']['Brioche']] == [svc.NOME_FORNO]
    # só MESA: só a etapa de mesa
    assert [e['nome'] for e in lido['fichas']['Sourdough Integral']] == \
        [svc.NOME_MESA]
    # local + MESA: só a mesa, com o aviso do local
    assert [e['nome'] for e in lido['fichas']['Pão de Cranberry']] == \
        [svc.NOME_MESA]
    so_local = [a for a in lido['avisos'] if 'só com o local' in a]
    assert sorted(a.split(':')[0] for a in so_local) == \
        ['Brioche', 'Pão de Cranberry']
    assert 'etapa não criada' in so_local[0]


def test_celula_amarela_em_branco_avisa_e_nao_vira_etapa(tmp_path):
    caminho = _planilha(tmp_path, {
        'caixa': [['Pain au Chocolat', None, None, (None, 'incerto'), None,
                   None, None, None]],
        'forno': [['Brioche', 170, (None, 'incerto'), 80, 100, None, None]],
    })
    lido = svc.ler_planilha(caminho)
    assert 'Pain au Chocolat' not in lido['fichas']       # linha vazia segue fora
    vazias = [a for a in lido['avisos'] if 'célula amarela em branco' in a]
    assert len(vazias) == 2
    assert any(a.startswith('Pain au Chocolat') and '(duração;' in a
               for a in vazias)
    assert any(a.startswith('Brioche') and '(tempo;' in a for a in vazias)
    forno = lido['fichas']['Brioche'][0]
    assert forno['duracao_min'] == 0
    assert 'incerto' not in forno['parametros']          # vazio não é valor


def test_validar_parametros_descarta_lixo_e_parametros_de_form_nunca_levanta():
    from app.services.etapas_receita import parametros_de_form, validar_parametros
    assert validar_parametros({'forno_c': 235, 'chave_estranha': 1,
                               'dobras_min': 5, 'alternativas': 7,
                               'teto': 'x', 'origem': 'a' * 500}) == {
        'forno_c': 235, 'origem': 'a' * 200}
    assert validar_parametros({'alternativas': [{'tempo_min': 'x'}, 'lixo',
                                                {'local': 'A', 'tempo_min': 360}]}) == {
        'alternativas': [{'tempo_min': 360, 'local': 'A'}]}
    assert validar_parametros({'incerto': ['forno_c', 'zzz', 3]}) == {'incerto': ['forno_c']}
    assert validar_parametros({'forno_c': float('inf')}) is None
    assert validar_parametros('x') is None and validar_parametros({}) is None
    # JSON aninhado dentro do teto não derruba o editor
    assert parametros_de_form('{"a":' + '[' * 1990 + ']' * 1990 + '}') is None
    assert parametros_de_form(json.dumps({'dobras_min': 5})) is None
    assert parametros_de_form(json.dumps({'forno_c': 235, 'lixo': 1})) == {'forno_c': 235}


def test_aplicar_recusa_quando_nenhuma_receita_casa_e_nada_apaga(app, lido):
    fora = _receita('Cookie', categoria='Doces')
    _etapa_velha(fora)
    with pytest.raises(ValueError, match='nada foi apagado'):
        svc.aplicar(lido['fichas'], svc.casar(lido['fichas']))
    assert ReceitaEtapa.query.filter_by(receita_id=fora.id).count() == 1
    assert svc.ultima_importacao() is None
    assert not [k for k in [c.key for c in AppConfig.query.all()]
                if k.startswith(svc.PREFIXO_BACKUP)]


def test_seed_sem_casamento_nao_grava_marker_nem_apaga(app):
    from app.migrations_legacy import _seed_fichas_producao_2026_10
    fora = _receita('Cookie', categoria='Doces')
    _etapa_velha(fora)
    _seed_fichas_producao_2026_10(app)
    assert AppConfig.get('seed_fichas_producao_2026_10') is None
    assert ReceitaEtapa.query.filter_by(receita_id=fora.id).count() == 1


def test_aplicar_pela_tela_sem_casamento_avisa(app, owner_user):
    import re
    fora = _receita('Cookie', categoria='Doces')
    _etapa_velha(fora)
    c = _login(app, owner_user)
    html = _upload(c).get_data(as_text=True)
    token = re.search(r'name="token" value="([^"]+)"', html).group(1)
    resp = c.post('/receitas/fichas-producao/aplicar', data={'token': token},
                  follow_redirects=True)
    assert 'nada foi apagado' in resp.get_data(as_text=True)
    assert ReceitaEtapa.query.filter_by(receita_id=fora.id).count() == 1


def test_aplicar_serializa_com_advisory_lock_em_postgres(app, lido, monkeypatch):
    _receita('Sourdough Tradicional')
    chamadas = []
    original = db.session.execute

    def _execute(stmt, *a, **kw):
        chamadas.append(str(stmt))
        return original(text('SELECT 1')) if 'pg_advisory_xact_lock' in str(stmt) \
            else original(stmt, *a, **kw)
    monkeypatch.setattr(db.engine.dialect, 'name', 'postgresql', raising=False)
    monkeypatch.setattr(db.session, 'execute', _execute)
    svc.aplicar(lido['fichas'], svc.casar(lido['fichas']))
    assert any('pg_advisory_xact_lock' in c for c in chamadas)
    assert svc.LOCK_KEY_APLICAR == 7768


def test_aplicar_a_toda_categoria_nao_copia_os_parametros(app, admin_user):
    sd = _receita('Sourdough Tradicional')
    si = _receita('Sourdough Integral')
    db.session.add(ReceitaEtapa(receita_id=sd.id, ordem=0, nome='Batimento',
                                duracao_min=15, equipamento='amassadeira',
                                parametros={'velocidade_1_min': 14,
                                            'origem': '1. Batimento · linha 5'}))
    db.session.add(ReceitaEtapa(receita_id=si.id, ordem=0, nome='Batimento',
                                duracao_min=26, equipamento='amassadeira',
                                parametros={'velocidade_1_min': 24}))
    db.session.commit()
    c = _login(app, admin_user)
    c.post(f'/receitas/{sd.id}/etapas', data={
        'acao': 'aplicar_categoria',
        'nome[]': ['Batimento'], 'duracao[]': ['15'], 'recurso[]': ['amassadeira'],
        'descricao[]': [''],
        'parametros[]': [json.dumps({'velocidade_1_min': 14,
                                     'origem': '1. Batimento · linha 5'})],
    }, follow_redirects=True)
    assert ReceitaEtapa.query.filter_by(receita_id=sd.id).first().parametros == {
        'velocidade_1_min': 14, 'origem': '1. Batimento · linha 5'}
    alvo = ReceitaEtapa.query.filter_by(receita_id=si.id).first()
    assert alvo.nome == 'Batimento' and alvo.duracao_min == 15
    assert alvo.parametros is None            # parâmetro é por receita


def test_round_trip_do_hidden_pelo_html_renderizado(app, admin_user):
    """O hidden parametros[] do editor carrega o JSON de verdade e re-POSTar
    o que a tela renderizou mantém os parâmetros iguais."""
    import re
    from html import unescape
    sd = _receita('Sourdough Tradicional')
    params = {'velocidade_1_min': 14, 'velocidade_2_min': 1, 'temp_liquido_c': 5.4,
              'origem': '1. Batimento · linha 5', 'incerto': ['temp_liquido_c']}
    db.session.add(ReceitaEtapa(receita_id=sd.id, ordem=0, nome='Batimento',
                                duracao_min=15, equipamento='amassadeira',
                                descricao='bata "devagar"', parametros=params))
    db.session.commit()
    c = _login(app, admin_user)
    html = c.get(f'/receitas/{sd.id}/etapas').get_data(as_text=True)
    hidden = re.findall(r'name="parametros\[\]" value="([^"]*)"', html)
    assert json.loads(unescape(hidden[0])) == params
    c.post(f'/receitas/{sd.id}/etapas', data={
        'nome[]': ['Batimento'], 'duracao[]': ['15'], 'recurso[]': ['amassadeira'],
        'descricao[]': ['bata "devagar"'], 'parametros[]': [unescape(hidden[0])],
    }, follow_redirects=True)
    e = ReceitaEtapa.query.filter_by(receita_id=sd.id).first()
    assert e.parametros == params and e.descricao == 'bata "devagar"'


def test_round_trip_do_hidden_na_ficha_do_padeiro(app):
    import re
    from html import unescape
    u = Usuario(nome='Padeiro RT', login='padrt', papel='padeiro')
    u.set_senha('12345678')
    db.session.add(u)
    sd = _receita('Sourdough Tradicional')
    params = {'forno_c': 235, 'teto': 100, 'lastro': 100}
    db.session.add(ReceitaEtapa(receita_id=sd.id, ordem=0, nome='Forno',
                                duracao_min=23, equipamento='forno', parametros=params))
    db.session.commit()
    c = app.test_client()
    c.post('/auth/login', data={'login': 'padrt', 'senha': '12345678'})
    html = c.get(f'/padeiro/fichas/{sd.id}').get_data(as_text=True)
    hidden = re.findall(r'name="parametros\[\]" value="([^"]*)"', html)
    assert json.loads(unescape(hidden[0])) == params
    c.post(f'/padeiro/fichas/{sd.id}', data={
        'nome[]': ['Forno'], 'duracao[]': ['25'], 'recurso[]': ['forno'],
        'descricao[]': ['asse com vapor'], 'parametros[]': [unescape(hidden[0])],
    })
    e = ReceitaEtapa.query.filter_by(receita_id=sd.id).first()
    assert e.parametros == params and e.duracao_min == 25


def test_token_da_previa_expira(app, owner_user, monkeypatch):
    import re

    from app.blueprints.receitas import routes as rotas
    _receita('Sourdough Tradicional')
    c = _login(app, owner_user)
    html = _upload(c).get_data(as_text=True)
    token = re.search(r'name="token" value="([^"]+)"', html).group(1)
    monkeypatch.setattr(rotas, '_FICHAS_TOKEN_MAX_AGE', -1)
    resp = c.post('/receitas/fichas-producao/aplicar', data={'token': token},
                  follow_redirects=True)
    assert 'prévia expirou' in resp.get_data(as_text=True)
    assert ReceitaEtapa.query.count() == 0


def test_previa_lista_ambigua_e_avisos(app, admin_user):
    _receita('Pão de Cranberry', categoria='Fornadas Especiais')
    _receita('Pão de Cranberry', categoria='Fornadas Especiais')
    c = _login(app, admin_user)
    html = _upload(c).get_data(as_text=True)
    assert 'ambígua' in html
    assert 'leitura incerta' in html and 'até a última dobra' in html


def test_ordem_aberta_de_receita_fora_da_planilha_fica_sem_processo(app, lido):
    _receita('Sourdough Tradicional')
    cookie = _receita('Cookie', categoria='Doces')
    _etapa_velha(cookie, 'Assar cookie')
    pl = PlanejamentoProducao(data=hoje(), origem='cronograma')
    db.session.add(pl)
    db.session.flush()
    it = PlanejamentoItem(planejamento_id=pl.id, receita_id=cookie.id,
                          multiplicador=1, qtd_alvo=10, produzido_qtd=0)
    db.session.add(it)
    db.session.flush()
    db.session.add(PlanejamentoItemBatelada(
        item_id=it.id, bateladas=1,
        dados={'unidades': 10, 'mp': [], 'processo': [{'nome': 'Assar cookie'}]}))
    db.session.commit()
    svc.aplicar(lido['fichas'], svc.casar(lido['fichas']))
    db.session.commit()
    snap = db.session.get(PlanejamentoItemBatelada, it.id)
    db.session.refresh(snap)
    assert snap.dados['processo'] == []        # o Gantt mostra "sem etapas"
    bk = json.loads(AppConfig.get(svc.ultima_importacao()['backup']))
    assert bk['etapas'][0]['nome'] == 'Assar cookie'
